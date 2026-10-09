"""Claims audit, count check and animation inventory for Paytitude_R2.pptx."""
import re, sys
from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn

deck = Presentation(sys.argv[1])


def all_text(shapes):
    out = []
    for sh in shapes:
        if sh.shape_type == 6:
            out += all_text(sh.shapes)
        elif sh.has_text_frame:
            out.append(sh.text_frame.text)
        elif getattr(sh, "has_table", False) and sh.has_table:
            out += [c.text for r in sh.table.rows for c in r.cells]
    return out


texts = []
for i, s in enumerate(deck.slides, 1):
    t = " ".join(all_text(s.shapes))
    n = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
    texts.append((i, t, n))

BANNED = [r"interviewed 50", r"we interviewed", r"50 merchants (said|told)", r"will switch", r"merchants will", r"\bguarantee",
          r"(?<!not )(?<!, not )\bforecast\b", r"\bapproved\b", r"only company", r"can't copy", r"cannot copy", r"escrow",
          r"is not credit", r"MDR (does not|doesn't) apply", r"2,700", r"₹27,000 crore of", r"cashback (offer|reward)s?\b(?! )",
          r"field research(?!\.)", r"pays suppliers directly", r"directly from collections", r"RBI (has )?approved", r"NPCI (has )?approved"]
ALLOWED_CONTEXT = {r"(?<!not )(?<!, not )\bforecast\b": ["not a forecast", "not forecast", "projection, not forecast", "PROJECTION, NOT FORECAST"],
                   r"field research(?!\.)": ["not field research"]}
print("== Claims audit (slide text + speaker notes)")
hits = 0
for pat in BANNED:
    for i, t, n in texts:
        for where, body in (("slide", t), ("notes", n)):
            for m in re.finditer(pat, body, flags=re.I):
                ctx = body[max(0, m.start() - 40): m.end() + 40].replace("\n", " ")
                if any(a.lower() in ctx.lower() for a in ALLOWED_CONTEXT.get(pat, [])):
                    continue
                hits += 1
                print(f"  HIT slide {i} [{where}] /{pat}/: …{ctx}…")
print(f"  {hits} unexplained hits")

print("\n== Mentions to review by hand (context should be negative or historical)")
for pat in [r"MDR", r"cashback", r"credit", r"ledger", r"rewards?"]:
    c = sum(len(re.findall(pat, t, flags=re.I)) for _, t, _ in texts)
    print(f"  /{pat}/ appears {c}x on slides")
    for i, t, _ in texts:
        for m in re.finditer(pat, t, flags=re.I):
            print(f"     s{i}: …{t[max(0, m.start() - 50): m.end() + 30]}…")

print("\n== Count check")
names = [(i, " ".join(all_text(s.shapes))[:40]) for i, s in enumerate(deck.slides, 1)]
div = [i for i, s in enumerate(deck.slides, 1) if any(sh.name == "Title" and sh.text_frame.text == "Annexures" for sh in s.shapes)]
print(f"  total slides {len(deck.slides)}; annexure divider at {div}; content slides between cover and divider: {div[0] - 2}")

print("\n== Animation and transition inventory")
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
for i, s in enumerate(deck.slides, 1):
    el = s._element
    tr = "none"
    for ch in el:
        tag = etree.QName(ch).localname
        if tag == "transition":
            tr = etree.QName(ch[0]).localname if len(ch) else "transition"
        elif tag == "AlternateContent":
            tr = "morph (fade fallback)"
    tim = el.find(qn("p:timing"))
    clicks = autos = effects = 0
    presets = set()
    if tim is not None:
        seq = tim.find(".//" + P + "cTn[@nodeType='mainSeq']")
        for par in seq.find(P + "childTnLst"):
            conds = par.find(P + "cTn").find(P + "stCondLst")
            if conds.find(P + "cond[@evt='onBegin']") is not None:
                autos += 1
            else:
                clicks += 1
        for c in tim.iter(P + "cTn"):
            if c.get("presetClass"):
                effects += 1
                presets.add(f"{c.get('presetClass')}:{c.get('presetID')}")
        # every spid must exist on the slide
        ids = {int(x.get("id")) for x in el.iter(qn("p:cNvPr"))}
        bad = [int(t.get("spid")) for t in tim.iter(P + "spTgt") if int(t.get("spid")) not in ids]
        assert not bad, (i, bad)
    print(f"  slide {i:2d}: transition={tr:22s} clicks={clicks} auto={autos} effects={effects} presets={sorted(presets)}")
