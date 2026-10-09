"""Check that every key number appears on its slide, and emit the audit table for build_report.md."""
import sys
from pptx import Presentation


def all_text(shapes):
    out = []
    for sh in shapes:
        if sh.shape_type == 6:
            out += all_text(sh.shapes)
        elif sh.has_text_frame:
            out.append(sh.text_frame.text)
        elif getattr(sh, "has_table", False) and sh.has_table:
            out += [c.text for r in sh.table.rows for c in r.cells]
    return out


deck = Presentation(sys.argv[1])
T = {i: " ".join(all_text(s.shapes)) for i, s in enumerate(deck.slides, 1)}
# deck index = content slide + 1 (cover is index 1); annexures A1..A8 = 18..25
ROWS = [
    (2, "1.57 crore", "S1 Paytm Q1 FY27 release", "Fact"),
    (2, "16 of 50, 7 yes / 9 maybe / 0 no", "Segment_Rules!C12:F12", "Synthetic"),
    (2, "+1.6 pts", "Share_Metric!C14 (1.57)", "Projection"),
    (2, "+3.9 pts", "Share_Metric!C13 (3.93)", "Projection"),
    (2, "₹2,660 crore", "Funnel_Model!C31 (2,661)", "Projection"),
    (2, "₹240–19,400 crore", "Funnel_Model!B31 (242), D31 (19,438)", "Projection"),
    (2, "₹3,200", "Unit_Costs!B24 (3,196)", "Projection"),
    (2, "3,000 vs 3,000", "Pilot_Design", "Proposal"),
    (3, "1.57 crore", "S1", "Fact"),
    (3, "8.0%", "S2 NPCI Aug 2026 (8.03%)", "Fact (secondary report)"),
    (3, "45.6%", "S2 (45.64%)", "Fact (secondary report)"),
    (3, "32.3%", "S2 (32.26%)", "Fact (secondary report)"),
    (3, "0 of 50", "VOC_Synthetic col I (no Paytm main app)", "Synthetic"),
    (3, "29 use PhonePe, 15 Google Pay", "VOC_Synthetic col I", "Synthetic"),
    (4, "41 of 50", "Segment_Rules!C33", "Synthetic"),
    (4, "PhonePe 29, Google Pay 15", "VOC_Synthetic col I", "Synthetic"),
    (5, "T+1", "S4 Paytm for Business support", "Fact"),
    (5, "₹1–3 a request, once an hour, not after 11 pm", "S6, S7 PhonePe help", "Fact"),
    (5, "₹12,400", "Illustrative amount", "Illustrative"),
    (6, "16 · Group 1", "Segment_Rules!C38", "Synthetic"),
    (6, "7 yes · 9 maybe · 0 no", "Segment_Rules!D38:F38", "Synthetic"),
    (6, "4 yes · 3 maybe · 2 no", "Segment_Rules!D40:F40", "Synthetic"),
    (6, "0 yes · 5 maybe · 2 no", "Segment_Rules!D39:F39", "Synthetic"),
    (6, "0 yes · 2 maybe · 6 no", "Segment_Rules!D41:F41", "Synthetic"),
    (6, "7 of 9 rows contradictory", "Segment_Rules row 40 vs contradiction flags", "Synthetic (audit)"),
    (6, "Group 3 (18)", "Segment_Rules!C14", "Synthetic"),
    (7, "₹40 lakh", "S10, S11 legal summaries of RBI PA Directions 2025", "Fact (primary text to confirm)"),
    (7, "₹6,800", "Illustrative amount", "Illustrative"),
    (8, "up to T+1", "S4", "Fact"),
    (9, "₹9,250", "Illustrative amount", "Illustrative"),
    (10, "₹1–3 a request", "S6", "Fact"),
    (11, "[__%]", "Validation_Log (empty)", "Placeholder"),
    (13, "40% of eligible", "Funnel_Model!C10", "Assumption"),
    (13, "15–25%", "Share_Metric early indicators", "Hypothesis"),
    (13, "50%", "Funnel_Model!C19", "Hypothesis"),
    (14, "≈3,000", "Pilot_Design", "Proposal"),
    (14, "+4 to +6 pts", "Share_Metric!C15 (4.6 before overlap)", "Hypothesis"),
    (14, "≥15%", "Pilot_Design gates", "Decision rule"),
    (14, "<8%", "Pilot_Design gates", "Decision rule"),
    (15, "+1.6 pts", "Share_Metric!C14", "Projection"),
    (15, "+3.9 pts", "Share_Metric!C13", "Projection"),
    (15, "₹2,660 crore", "Funnel_Model!C31", "Projection"),
    (15, "₹1.7 lakh crore", "Share_Metric!C10 (1.69)", "Projection (addressable pool)"),
    (15, "₹3,900 with build", "Unit_Costs!B25 (3,903)", "Projection"),
    (15, "₹0.009 per ₹", "Unit_Costs!B26 (0.0093)", "Projection"),
    (15, "₹0.011 with build", "Unit_Costs!B27 (0.0114)", "Projection"),
    (15, "≈65 MTU per 1,000", "Funnel_Model!C32 (65.0)", "Projection"),
    (15, "≈₹0.023", "Unit_Costs!B30 (0.0228)", "Projection"),
    (16, "3,000 vs 3,000", "Pilot_Design", "Proposal"),
    (16, "25%", "Solution doc §13 P0 hypothesis threshold (G0)", "Decision rule"),
    (22, "20.9 L", "Funnel_Model!B9", "Projection"),
    (22, "29.9 L", "Funnel_Model!C9", "Projection"),
    (22, "35.9 L", "Funnel_Model!D9", "Projection"),
    (22, "12.0 L", "Funnel_Model!C11", "Projection"),
    (22, "23.1%", "Funnel_Model!C17", "Assumption × synthetic"),
    (22, "2.77 L", "Funnel_Model!C18", "Projection"),
    (22, "1.38 L", "Funnel_Model!C20", "Projection"),
    (22, "78,000", "Funnel_Model!C22 (77,799)", "Projection"),
    (22, "10.4 L", "Funnel_Model!C26", "Projection"),
    (22, "₹222 cr", "Funnel_Model!C30 (221.8)", "Projection"),
    (22, "₹14,100 crore/month", "Share_Metric!C9 (14,102)", "Projection (addressable)"),
    (22, "₹0.75", "Unit_Costs!B11 (0.753)", "Assumption"),
    (22, "₹3,080", "Unit_Costs!B31 (3,075)", "Projection"),
    (22, "₹2,400–4,700", "Unit_Costs!B32:B33 (2,428 / 4,734)", "Projection"),
    (24, "≈₹27,000 crore", "Round 1 PDF (S17)", "Round 1 figure (dropped)"),
    (24, "500 Andheri merchants", "Round 1 PDF (S17)", "Round 1 figure"),
]
bad = 0
lines = ["| Slide | Number on slide | Source (cell or reference) | Status | Present |", "| --- | --- | --- | --- | --- |"]
for idx, num, src, status in ROWS:
    ok = num in T[idx]
    bad += not ok
    label = "Cover" if idx == 1 and False else (str(idx - 1) if 2 <= idx + 0 <= 16 or idx == 1 else f"A{idx - 17}")
    label = str(idx - 1) if 2 <= idx <= 16 else ("1" if idx == 1 else f"A{idx - 17}")
    lines.append(f"| {label if idx != 1 else '1'} | {num} | {src} | {status} | {'yes' if ok else '**MISSING**'} |")
print("\n".join(lines))
print(f"\n{len(ROWS) - bad}/{len(ROWS)} numbers found on their slides", file=sys.stderr)
