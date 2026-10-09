"""Recompute slide 10 (solution validation) from the annexure workbook's Validation_Log and re-export the PDF.

Usage:
    python3 fill_validation.py Paytitude_R2_Annexure.xlsx Paytitude_R2.pptx [--out Paytitude_R2.pptx]
                               [--pdf Paytitude_R2.pdf] [--exploratory-n 5]

Reads one row per tested merchant from row 19 of 'Validation_Log' (rows with an ID in column A), computes the
behaviour metrics, writes them into the magenta placeholders on slide 10 (and the matching line in the executive
summary and annexure A1), recolours them, and re-exports the PDF with LibreOffice. It never invents values:
an empty log leaves the deck untouched.
"""
import argparse
import os
import shutil
import subprocess
import sys
import tempfile
from collections import Counter

import openpyxl
from pptx import Presentation
from pptx.dml.color import RGBColor

MAGENTA = RGBColor.from_string("D6249F")
NAVY = RGBColor.from_string("002E6E")
CYAN_D = RGBColor.from_string("0077B6")
FIRST_ROW, LAST_ROW = 19, 78
COLS = ["id", "date", "city", "type", "collector", "settlement", "vendor_upi_week", "app", "method", "task",
        "commitment_text", "named", "trial", "enabled", "used_before", "objection", "verbatim", "evidence", "of_five"]


def read_log(path):
    ws = openpyxl.load_workbook(path, data_only=True)["Validation_Log"]
    rows = []
    for r in range(FIRST_ROW, LAST_ROW + 1):
        vals = [ws.cell(row=r, column=c).value for c in range(1, len(COLS) + 1)]
        if vals[0] not in (None, ""):
            rows.append(dict(zip(COLS, vals)))
    return rows


def yes(v):
    return str(v or "").strip().upper() in ("Y", "YES")


def metrics(rows):
    n = len(rows)
    pct = lambda k: round(100 * sum(yes(r[k]) for r in rows) / n)
    anyc = round(100 * sum(any(yes(r[k]) for k in ("named", "trial", "enabled")) for r in rows) / n)
    sett = Counter(str(r["settlement"] or "").strip().lower() for r in rows)
    obj = Counter(str(r["objection"] or "").strip().lower() for r in rows if r["objection"])
    methods = Counter(str(r["method"] or "").strip() for r in rows if r["method"])
    return dict(n=n, task=pct("task"), task_n=sum(yes(r["task"]) for r in rows), named=pct("named"), trial=pct("trial"),
                enabled=pct("enabled"), any=anyc, used_before=pct("used_before"),
                same=sett.get("same day", 0), next=sett.get("next day", 0), sep=sett.get("separate account", 0),
                obj=obj, methods=methods)


def iter_shapes(shapes):
    for sh in shapes:
        if sh.shape_type == 6:
            yield from iter_shapes(sh.shapes)
        else:
            yield sh


def magenta_runs(shape):
    out = []
    if not shape.has_text_frame:
        return out
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            try:
                if r.font.color and r.font.color.type is not None and r.font.color.rgb == MAGENTA:
                    out.append(r)
            except AttributeError:
                pass
    return out


def set_run(run, text, color=NAVY):
    run.text = text
    run.font.color.rgb = color


def fill(prs, m, exploratory_n=None):
    val_slide = next(s for s in prs.slides if any(sh.name == "VAL_title" for sh in iter_shapes(s.shapes)))
    by_name = {sh.name: sh for sh in iter_shapes(val_slide.shapes)}
    # headline: an assertion built from the actual result (edit wording if needed, never the numbers)
    title = by_name["VAL_title"]
    head = (f"{m['task_n']} of {m['n']} target merchants completed the Settle & Pay task unaided, and "
            f"{m['named']}% named a supplier payment they would move to it")
    p = title.text_frame.paragraphs[0]
    for extra in p.runs[1:]:
        extra.text = ""
    set_run(p.runs[0], head)
    for sh in iter_shapes(val_slide.shapes):
        if sh.name == "Eyebrow":
            sh.text_frame.paragraphs[0].runs[0].text = "SOLUTION VALIDATION"
    set_run(magenta_runs(by_name["VAL_n"])[0], f"n = {m['n']} ")
    for nm, key in [("VAL_t1", "task"), ("VAL_t2", "named"), ("VAL_t3", "trial"), ("VAL_t4", "enabled")]:
        set_run(magenta_runs(by_name[nm])[0], f"{m[key]}%", CYAN_D)
        by_name[nm].line.color.rgb = CYAN_D
        by_name[nm].line.dash_style = None
    by_name["VAL_n"].line.color.rgb = CYAN_D
    by_name["VAL_n"].line.dash_style = None
    top = ", ".join(f"{k.title()} {v}" for k, v in m["obj"].most_common(4)) or "none recorded"
    ctx = magenta_runs(by_name["VAL_ctx"])
    set_run(ctx[0], f"{m['used_before']}%")
    set_run(ctx[1], f"{m['same']} same day / {m['next']} next day / {m['sep']} separate account")
    set_run(ctx[2], top)
    if exploratory_n is not None:
        for r in magenta_runs(by_name["VAL_context_evidence"]):
            set_run(r, f"n = {exploratory_n}")
    # executive summary band and annexure A1 strip
    for s in prs.slides:
        for sh in iter_shapes(s.shapes):
            for r in magenta_runs(sh):
                if "results pending" in r.text:
                    set_run(r, f"n = {m['n']}; {m['task']}% completed the task, {m['any']}% made a commitment")
                elif r.text.startswith("VALIDATION (≥20"):
                    set_run(r, "VALIDATION (≥20 TARGET USERS): ")
                elif r.text == "n = [__]" and exploratory_n is not None:
                    set_run(r, f"n = {exploratory_n}")
            if sh.has_text_frame and "not yet run at the time of this build" in sh.text_frame.text:
                for p_ in sh.text_frame.paragraphs:
                    for r in p_.runs:
                        if "not yet run" in r.text:
                            r.text = f"run with n = {m['n']} target merchants; results on slide 10, rows in 'Validation_Log'."
    for s in prs.slides:
        for g in s.shapes:
            if g.shape_type == 6 and g.name == "band3":
                for sh in g.shapes:
                    if sh.has_text_frame and sh.text_frame.text.strip() == "TO VALIDATE":
                        sh.text_frame.paragraphs[0].runs[0].text = f"VALIDATED (n = {m['n']})"
                        sh.fill.solid(); sh.fill.fore_color.rgb = CYAN_D
                        sh.width = int(sh.width * 1.55); sh.left = int(12.73 * 914400) - sh.width
    return head


def export_pdf(pptx_path, pdf_path):
    with tempfile.TemporaryDirectory() as d:
        shutil.copy(pptx_path, os.path.join(d, "deck.pptx"))
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", d, os.path.join(d, "deck.pptx")],
                       check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=300)
        shutil.copy(os.path.join(d, "deck.pdf"), pdf_path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx"); ap.add_argument("pptx")
    ap.add_argument("--out"); ap.add_argument("--pdf"); ap.add_argument("--exploratory-n", type=int)
    a = ap.parse_args()
    rows = read_log(a.xlsx)
    if not rows:
        sys.exit("Validation_Log has no rows yet (fill from row 19). Deck left unchanged.")
    m = metrics(rows)
    if m["n"] < 20:
        print(f"WARNING: only {m['n']} merchants logged; Round 2 requires at least 20.")
    prs = Presentation(a.pptx)
    head = fill(prs, m, a.exploratory_n)
    out = a.out or a.pptx
    prs.save(out)
    print("Slide 10 headline:", head)
    print({k: v for k, v in m.items() if k not in ("obj", "methods")}, dict(m["obj"]), dict(m["methods"]))
    if a.pdf:
        export_pdf(out, a.pdf)
        print("PDF re-exported:", a.pdf)
    print("Now rewrite the slide 10 speaker notes and check the headline wording against the evidence.")


if __name__ == "__main__":
    main()
