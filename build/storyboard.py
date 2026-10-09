"""Write storyboard.md from the deck builder's registry (so storyboard and deck never drift)."""
from pptx.util import Emu

FOOTER_NAMES = {"FooterTag", "Sources", "SlideNo"}

COMPLIANCE = [
    ("General", "Build on Round 1; same track and segment; say what changed and why", "1 (change band), 2, 5, A7"),
    ("Deliverables", "15-slide deck excluding cover and annexures", "Slides 1–15 between cover and 'Annexures' divider"),
    ("Deliverables", "Opens with a one-slide executive summary", "1"),
    ("Deliverables", "Optional annexures: prototypes, extra VOCs", "8, A2, A3; workbook VOC_Synthetic, VOC_Real"),
    ("Deliverables", "File names TeamName_R2.pdf and TeamName_R2_Annexure.xlsx", "Paytitude_R2.pdf, Paytitude_R2_Annexure.xlsx"),
    ("1. Validation", "Test with ≥20 target users (concept test, prototype walkthrough or trial)", "10 (template, results pending), A8, Validation_Log"),
    ("1. Validation", "Report what users did or committed to", "10 (behaviour tiles: observed / commitment / action)"),
    ("1. Validation", "State method, n, secondary sources, limits", "10, A1"),
    ("2. Design & feasibility", "User journey before and after", "4 (before), 7 (before + after), 8 (concept screens), A3"),
    ("2. Design & feasibility", "Paytm features used; what is new", "6, 11"),
    ("2. Design & feasibility", "Regulatory or partner dependencies", "6, 11, A6"),
    ("2. Design & feasibility", "Rough build effort (S / M / L)", "11"),
    ("3. GTM & pilot", "Target segment", "5, 12"),
    ("3. GTM & pilot", "Channels, partners, messaging", "12"),
    ("3. GTM & pilot", "Pilot cohort and control group", "13, A4"),
    ("3. GTM & pilot", "Baseline and targets", "13, 14, A4"),
    ("3. GTM & pilot", "What result means scale up or stop", "13, 15"),
    ("3. GTM & pilot", "How to scale afterwards", "13, 15"),
    ("4. Business impact", "Main metric: increase in Paytm share of segment UPI GMV", "14, A5, Share_Metric"),
    ("4. Business impact", "Early indicators (activation, repeat)", "14, 12, A4"),
    ("4. Business impact", "Size of the opportunity", "14, A5"),
    ("4. Business impact", "Cost per extra active user and per ₹ of extra GMV", "14, A5, Unit_Costs"),
    ("4. Business impact", "Whether usage holds once incentives stop; no cashback dependence", "9, 14, 13"),
    ("4. Business impact", "Sensitivity on 2–3 key assumptions", "14, Sensitivity"),
    ("4. Business impact", "Key risks incl. competitor response, and mitigation", "9, 15, A6"),
]


def _texts(shapes):
    out = []
    for sh in shapes:
        if sh.shape_type == 6:  # group
            out += _texts(sh.shapes)
        elif sh.has_text_frame and sh.name not in FOOTER_NAMES:
            t = " ".join(p.text for p in sh.text_frame.paragraphs if p.text.strip())
            if t.strip():
                out.append((sh, t.strip()))
        elif getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                out.append((sh, " | ".join(c.text for c in row.cells)))
    return out


def _region(sh):
    e = lambda v: f"{Emu(v).inches:.2f}"
    return f"({e(sh.left)}, {e(sh.top)}, {e(sh.width)}×{e(sh.height)})"


def write(SB, path, N):
    L = ["# Storyboard: Paytitude R2 (Settle & Pay)", "",
         "Generated from the deck builder (`build/build_deck.py`), so the text below is exactly what is on each slide. "
         "Canvas 13.333 × 7.5 in; coordinates are (x, y, w×h) in inches.", "",
         "## Headlines alone (the argument in sequence)", ""]
    for d in SB:
        if isinstance(d["num"], int):
            L.append(f"{d['num']}. {d['headline']}")
    L += ["", "## Round 2 compliance matrix", "", "| Section | Guideline item | Where it is met |", "| --- | --- | --- |"]
    L += [f"| {a} | {b} | {c} |" for a, b, c in COMPLIANCE]
    L += ["", "## Slide by slide", ""]
    for d in SB:
        s = d["slide"]
        items = _texts(s.shapes)
        words = sum(len(t.split()) for _, t in items)
        L += [f"### {d['num']} · {d['role']}", "",
              f"**Headline:** {d['headline']}", "",
              f"**Core message:** {d['message']}", "",
              f"**Guideline served:** {d['guideline'] or '—'}", "",
              f"**Visual:** {d['visual']}", "",
              f"**Status chips:** {d['chips'] if isinstance(d['chips'], str) else ', '.join(d['chips']) or '—'}  ·  **Sources:** {d['sources']}", "",
              f"**On-slide text** ({words} words excluding footer):", ""]
        for sh, t in items:
            L.append(f"- `{sh.name}` {_region(sh)}: {t}")
        L += ["", "**Animation:**", ""]
        if d["anim"]:
            for k, (mode, lab) in enumerate(d["anim"], 1):
                L.append(f"{k}. {mode}: {lab}")
        else:
            L.append("None (static).")
        L += ["", "**Speaker notes:**", "", "> " + d["notes"], ""]
    open(path, "w").write("\n".join(L) + "\n")
    print("storyboard written", path)
