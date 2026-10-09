"""Drawing, animation and transition helpers for the Paytitude R2 deck (python-pptx + lxml)."""
from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

# ------------------------------------------------------------------ palette
NAVY = "002E6E"
NAVY2 = "0B3F86"      # raised surface on dark slides
NAVY3 = "1D4F96"
CYAN = "00BAF2"
CYAN_D = "0077B6"     # cyan family, readable as text on white
PALE = "EEF5FC"
PALE2 = "DCEBF8"
GREY = "5B6B7C"
GREY_L = "C9D3DE"
GREY_XL = "F3F5F8"
INK = "1B2733"
WHITE = "FFFFFF"
SOFT = "A9C1E0"       # secondary text on dark
RED = "C62828"
GREEN = "1E8E3E"
AMBER = "F2B200"
BLUE = "2F6FD6"
VGREY = "6B7785"
MAGENTA = "D6249F"    # placeholder colour: to be filled from validation data
FONT = "Calibri"

CHIP = {
    "fact": ("FACT", GREEN, WHITE),
    "synthetic": ("SYNTHETIC", AMBER, INK),
    "projection": ("PROJECTION", BLUE, WHITE),
    "validate": ("TO VALIDATE", VGREY, WHITE),
    "validated": ("VALIDATED (n = __)", CYAN_D, WHITE),
    "placeholder": ("PLACEHOLDER", MAGENTA, WHITE),
    "concept": ("CONCEPT PROTOTYPE", NAVY, WHITE),
    "genuine": ("GENUINE · EXPLORATORY", CYAN_D, WHITE),
}


def rgb(h):
    return RGBColor.from_string(h)


def new_presentation():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    return prs


def blank(prs, dark=False):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fill = s.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(NAVY if dark else WHITE)
    return s


# ------------------------------------------------------------------ text
def _apply_run(run, size, bold, color, italic=False, font=FONT):
    f = run.font
    f.name, f.size, f.bold, f.italic = font, Pt(size), bold, italic
    f.color.rgb = rgb(color)


def text(container, x, y, w, h, paras, size=12, bold=False, color=INK, align="l", anchor="t",
         name=None, fill=None, line=None, margin=0.04, spacing=None, line_spacing=None, italic=False,
         shape=None, radius=None, line_w=0.75, dash=False):
    """Add a text box (or an autoshape with text if fill/line/shape given).

    paras: str | list of paragraphs; each paragraph is str or list of runs;
           a run is str or (text, dict(size=, bold=, color=, italic=)).
    """
    if shape or fill or line:
        shp = container.shapes.add_shape(shape or MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        if radius is not None:
            shp.adjustments[0] = radius
        if fill:
            shp.fill.solid(); shp.fill.fore_color.rgb = rgb(fill)
        else:
            shp.fill.background()
        if line:
            shp.line.color.rgb = rgb(line); shp.line.width = Pt(line_w)
            if dash:
                shp.line.dash_style = 4  # dash
        else:
            shp.line.fill.background()
        shp.shadow.inherit = False
    else:
        shp = container.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        shp.name = name
    tf = shp.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(paras, str):
        paras = [paras]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        if spacing is not None and i > 0:
            para.space_before = Pt(spacing)
        if line_spacing:
            para.line_spacing = line_spacing
        runs = p if isinstance(p, list) else [p]
        for r in runs:
            t, o = (r, {}) if isinstance(r, str) else r
            run = para.add_run()
            run.text = t
            _apply_run(run, o.get("size", size), o.get("bold", bold), o.get("color", color), o.get("italic", italic))
    return shp


def box(container, x, y, w, h, fill=None, line=None, name=None, shape=MSO_SHAPE.RECTANGLE, radius=None,
        line_w=0.75, dash=False):
    shp = container.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius is not None:
        shp.adjustments[0] = radius
    if fill:
        shp.fill.solid(); shp.fill.fore_color.rgb = rgb(fill)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = rgb(line); shp.line.width = Pt(line_w)
        if dash:
            shp.line.dash_style = 4
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    if name:
        shp.name = name
    if shp.has_text_frame:
        shp.text_frame.text = ""
    return shp


def line(container, x1, y1, x2, y2, color=GREY_L, width=1.25, arrow=False, name=None, dash=False):
    c = container.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    if dash:
        c.line.dash_style = 4
    if arrow:
        ln = c.line._get_or_add_ln()
        tail = etree.SubElement(ln, qn("a:tailEnd"))
        tail.set("type", "triangle"); tail.set("w", "med"); tail.set("len", "med")
    if name:
        c.name = name
    return c


def group(slide, name):
    g = slide.shapes.add_group_shape()
    g.name = name
    return g


_FONTS = {}


def text_width(t, size, bold=False):
    """Width in inches of a single line in Calibri (measured with metric-compatible Carlito)."""
    from PIL import ImageFont
    key = bold
    if key not in _FONTS:
        _FONTS[key] = ImageFont.truetype(f"/usr/share/fonts/truetype/crosextra/Carlito-{'Bold' if bold else 'Regular'}.ttf", 200)
    return _FONTS[key].getlength(t) / 200 * size / 72


def chip(container, x, y, kind, label=None, h=0.22, size=9, name=None):
    lab, bg, fg = CHIP[kind]
    lab = label or lab
    w = chip_w(kind, label, size)
    return text(container, x, y, w, h, [[(lab, dict(bold=True, color=fg, size=size))]], align="c", anchor="m",
                fill=bg, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5, margin=0.02, name=name)


def chip_w(kind, label=None, size=9):
    lab = label or CHIP[kind][0]
    return text_width(lab, size, True) + 0.22


# ------------------------------------------------------------------ slide frame
def frame(slide, num, eyebrow, title, dark=False, sources="", title_size=26, title_color=None):
    fg = WHITE if dark else NAVY
    eb = CYAN if dark else GREY
    text(slide, 0.6, 0.3, 9.5, 0.28, [[(eyebrow.upper(), dict(bold=True, color=eb, size=11))]], name="Eyebrow")
    text(slide, 0.6, 0.6, 12.13, 1.1, title, size=title_size, bold=True, color=title_color or fg,
         name="Title", line_spacing=0.95)
    foot = SOFT if dark else GREY
    text(slide, 0.6, 7.0, 1.7, 0.3, "Track B · Paytitude", size=9, bold=True, color=foot, name="FooterTag", anchor="t")
    if sources:
        text(slide, 2.35, 6.98, 10.0, 0.42, sources, size=9, color=foot, name="Sources", anchor="t")
    if num is not None:
        text(slide, 12.33, 7.0, 0.4, 0.3, str(num), size=9, bold=True, color=foot, align="r", name="SlideNo")


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


# ------------------------------------------------------------------ phone mock-up
def phone(container, x, y, w, h, header, header_fill=NAVY, header_color=WHITE, name_prefix="phone", dark_bg=False):
    """Draw a phone body and return (inner_x, inner_y, inner_w) for content placement."""
    body = box(container, x, y, w, h, fill="1A2230", shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.09,
               name=f"{name_prefix}_body")
    m = 0.09
    box(container, x + m, y + m * 2.2, w - 2 * m, h - m * 4.2, fill=WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
        radius=0.04, name=f"{name_prefix}_screen")
    box(container, x + w / 2 - 0.25, y + 0.07, 0.5, 0.06, fill="3A4558", shape=MSO_SHAPE.ROUNDED_RECTANGLE,
        radius=0.5, name=f"{name_prefix}_speaker")
    text(container, x + m, y + m * 2.2, w - 2 * m, 0.42, [[(header, dict(bold=True, color=header_color, size=11))]],
         fill=header_fill, anchor="m", margin=0.1, name=f"{name_prefix}_header")
    return x + m + 0.12, y + m * 2.2 + 0.52, w - 2 * m - 0.24


# ------------------------------------------------------------------ charts
def bar_chart(slide, x, y, w, h, cats, vals, colors, number_format, name, label_size=12, cat_size=11,
              max_val=None, label_color=INK, gap=45):
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series("Series 1", vals)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    gf.name = name
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    ch.font.name = FONT
    ch.font.size = Pt(cat_size)
    plot = ch.plots[0]
    plot.gap_width = gap
    plot.vary_by_categories = False
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = number_format
    dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(label_size); dl.font.bold = True; dl.font.name = FONT
    dl.font.color.rgb = rgb(label_color)
    ser = plot.series[0]
    for i, c in enumerate(colors):
        pt = ser.points[i]
        pt.format.fill.solid(); pt.format.fill.fore_color.rgb = rgb(c)
        pt.format.line.fill.background()
    ca = ch.category_axis
    ca.reverse_order = True
    ca.has_major_gridlines = False
    ca.format.line.color.rgb = rgb(GREY_L)
    ca.tick_labels.font.size = Pt(cat_size); ca.tick_labels.font.name = FONT
    ca.tick_labels.font.color.rgb = rgb(INK)
    va = ch.value_axis
    va.visible = False
    va.has_major_gridlines = False
    va.minimum_scale = 0
    if max_val:
        va.maximum_scale = max_val
    return gf


# ------------------------------------------------------------------ native table
def native_table(slide, x, y, w, col_w, rows, header_fill=NAVY, size=12, row_h=0.32, name="Table",
                 zebra=True, bold_first_col=False, header_size=None, header_h=0.36):
    nrows, ncols = len(rows), len(rows[0])
    gf = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(w), Inches(row_h * nrows))
    gf.name = name
    tbl = gf.table
    tbl.first_row = True
    tbl.horz_banding = False
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    for i in range(nrows):
        tbl.rows[i].height = Inches(header_h if i == 0 else row_h)
        for j in range(ncols):
            cell = tbl.cell(i, j)
            val = rows[i][j]
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.fill.solid()
            if i == 0:
                cell.fill.fore_color.rgb = rgb(header_fill)
            else:
                cell.fill.fore_color.rgb = rgb(PALE if (zebra and i % 2 == 0) else WHITE)
            tf = cell.text_frame
            tf.word_wrap = True
            tf.paragraphs[0].text = ""
            runs = val if isinstance(val, list) else [val]
            p = tf.paragraphs[0]
            for r in runs:
                t, o = (r, {}) if isinstance(r, str) else r
                run = p.add_run(); run.text = str(t)
                _apply_run(run, o.get("size", header_size or size if i == 0 else size),
                           o.get("bold", i == 0 or (bold_first_col and j == 0)),
                           o.get("color", WHITE if i == 0 else INK))
    return gf


# ------------------------------------------------------------------ animation
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
WIPE = {"left": ("8", "wipe(left)"), "right": ("2", "wipe(right)"), "up": ("1", "wipe(up)"), "down": ("4", "wipe(down)")}


def fx(shape, effect="fade", trigger="with", delay=0, dur=None):
    """One effect. trigger: 'with' (same time as previous) or 'after' (after previous ends)."""
    return dict(shape=shape, effect=effect, trigger=trigger, delay=delay, dur=dur)


class Anim:
    def __init__(self, slide):
        self.slide = slide
        self.groups = []   # list of (mode, [fx]) where mode in {"click", "auto"}
        self.log = []      # human-readable description for the storyboard

    def click(self, effects, label=""):
        self.groups.append(("click", effects)); self.log.append(("Click", label))
        return self

    def auto(self, effects, label=""):
        self.groups.append(("auto", effects)); self.log.append(("Auto", label))
        return self

    def _effect_xml(self, e, ctn_id, node_type):
        spid = e["shape"].shape_id
        eff = e["effect"]
        ids = iter(range(ctn_id + 1, ctn_id + 10))
        if eff == "pulse":
            dur = e["dur"] or 300
            inner = (f'<p:animScale><p:cBhvr><p:cTn id="{next(ids)}" dur="{dur}" autoRev="1" fill="hold"/>'
                     f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr><p:by x="108000" y="108000"/></p:animScale>')
            head = f'presetID="6" presetClass="emph" presetSubtype="0" grpId="1"'
            return (f'<p:par><p:cTn id="{ctn_id}" {head} fill="hold" nodeType="{node_type}">'
                    f'<p:stCondLst><p:cond delay="{e["delay"]}"/></p:stCondLst><p:childTnLst>{inner}</p:childTnLst></p:cTn></p:par>'), 2 * dur, ctn_id + 2, 1
        dur = e["dur"] or (550 if eff.startswith("wipe") else 450)
        if eff == "fade":
            preset, sub, filt = "10", "0", "fade"
        elif eff.startswith("wipe"):
            preset = "22"
            sub, filt = WIPE[eff.split("-")[1]]
        else:
            raise ValueError(eff)
        set_id, eff_id = next(ids), next(ids)
        inner = (f'<p:set><p:cBhvr><p:cTn id="{set_id}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
                 f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst>'
                 f'</p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
                 f'<p:animEffect transition="in" filter="{filt}"><p:cBhvr><p:cTn id="{eff_id}" dur="{dur}"/>'
                 f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>')
        head = f'presetID="{preset}" presetClass="entr" presetSubtype="{sub}" grpId="0"'
        return (f'<p:par><p:cTn id="{ctn_id}" {head} fill="hold" nodeType="{node_type}">'
                f'<p:stCondLst><p:cond delay="{e["delay"]}"/></p:stCondLst><p:childTnLst>{inner}</p:childTnLst></p:cTn></p:par>'), dur, ctn_id + 3, 0

    def write(self):
        if not self.groups:
            return
        nid = 3
        pars = []
        builds = {}
        for gi, (mode, effects) in enumerate(self.groups):
            outer_id = nid; nid += 1
            if mode == "auto":
                st = '<p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond>'
            else:
                st = '<p:cond delay="indefinite"/>'
            inner_pars = []   # list of [pid, start, [xml...], max_end]
            for k, e in enumerate(effects):
                if k == 0:
                    node = "clickEffect" if mode == "click" else ("afterEffect" if e["trigger"] == "after" else "withEffect")
                    inner_pars.append([nid, 0, [], 0]); nid += 1
                elif e["trigger"] == "after":
                    node = "afterEffect"
                    prev = inner_pars[-1]
                    inner_pars.append([nid, prev[1] + prev[3], [], 0]); nid += 1
                else:
                    node = "withEffect"
                cur = inner_pars[-1]
                xml, length, nid, grp = self._effect_xml(e, nid, node)
                cur[2].append(xml)
                cur[3] = max(cur[3], e["delay"] + length)
                builds[(e["shape"].shape_id, grp)] = e["shape"]
            inner_xml = ""
            for pid, start, xs, _ in inner_pars:
                inner_xml += (f'<p:par><p:cTn id="{pid}" fill="hold"><p:stCondLst><p:cond delay="{start}"/></p:stCondLst>'
                              f'<p:childTnLst>{"".join(xs)}</p:childTnLst></p:cTn></p:par>')
            pars.append(f'<p:par><p:cTn id="{outer_id}" fill="hold"><p:stCondLst>{st}</p:stCondLst>'
                        f'<p:childTnLst>{inner_xml}</p:childTnLst></p:cTn></p:par>')
        bld = ""
        for (sid, grp), shp in sorted(builds.items(), key=lambda t: t[0]):
            tag = etree.QName(shp._element).localname
            if tag == "sp" and shp._element.find(qn("p:txBody")) is not None:
                bld += f'<p:bldP spid="{shp.shape_id}" grpId="{grp}" animBg="1"/>'
            elif tag == "graphicFrame":
                bld += f'<p:bldGraphic spid="{shp.shape_id}" grpId="{grp}"><p:bldAsOne/></p:bldGraphic>'
        xml = (f'<p:timing xmlns:p="{P_NS}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">'
               f'<p:childTnLst><p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq">'
               f'<p:childTnLst>{"".join(pars)}</p:childTnLst></p:cTn>'
               f'<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
               f'<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>'
               f'</p:seq></p:childTnLst></p:cTn></p:par></p:tnLst>'
               + (f'<p:bldLst>{bld}</p:bldLst>' if bld else "") + '</p:timing>')
        el = etree.fromstring(xml)
        sld = self.slide._element
        old = sld.find(qn("p:timing"))
        if old is not None:
            sld.remove(old)
        _insert_after_transition(sld, el)


def _insert_after_transition(sld, el):
    """Insert <p:timing> after <p:transition>/AlternateContent, else after clrMapOvr."""
    anchor = None
    for child in sld:
        tag = etree.QName(child).localname
        if tag in ("cSld", "clrMapOvr", "transition", "AlternateContent"):
            anchor = child
    anchor.addnext(el)


def transition(slide, kind="fade"):
    sld = slide._element
    for child in list(sld):
        if etree.QName(child).localname in ("transition", "AlternateContent"):
            sld.remove(child)
    if kind == "fade":
        el = etree.fromstring(f'<p:transition xmlns:p="{P_NS}" spd="med"><p:fade/></p:transition>')
    elif kind == "morph":
        el = etree.fromstring(
            '<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
            '<mc:Choice xmlns:p159="http://schemas.microsoft.com/office/powerpoint/2015/09/main" Requires="p159">'
            f'<p:transition xmlns:p="{P_NS}" xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main" spd="slow" p14:dur="1000">'
            '<p159:morph option="byObject"/></p:transition></mc:Choice>'
            f'<mc:Fallback><p:transition xmlns:p="{P_NS}" spd="slow"><p:fade/></p:transition></mc:Fallback>'
            '</mc:AlternateContent>')
    else:
        raise ValueError(kind)
    clr = sld.find(qn("p:clrMapOvr"))
    (clr if clr is not None else sld.find(qn("p:cSld"))).addnext(el)


def strip_theme_effects(prs):
    """Point every shape style's effectRef at idx 0 so no renderer adds theme shadows."""
    for slide in prs.slides:
        for el in slide._element.iter(qn("a:effectRef")):
            el.set("idx", "0")
