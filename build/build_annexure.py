"""Build Paytitude_R2_Annexure.xlsx with live formulas.

Usage: python3 build_annexure.py <voc_workbook.xlsx> <out.xlsx>
"""
import sys
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

import content as C

VOC_SRC, OUT = sys.argv[1], sys.argv[2]

NAVY, CYAN, PALE, GREY = "002E6E", "00BAF2", "EEF5FC", "5B6B7C"
FONT = "Arial"
f_base = Font(name=FONT, size=10)
f_bold = Font(name=FONT, size=10, bold=True)
f_head = Font(name=FONT, size=10, bold=True, color="FFFFFF")
f_title = Font(name=FONT, size=14, bold=True, color=NAVY)
f_note = Font(name=FONT, size=9, italic=True, color=GREY)
f_input = Font(name=FONT, size=10, color="0000FF")
f_link = Font(name=FONT, size=10, color="008000")
f_banner = Font(name=FONT, size=12, bold=True, color="FFFFFF")
fill_head = PatternFill("solid", fgColor=NAVY)
fill_pale = PatternFill("solid", fgColor=PALE)
fill_key = PatternFill("solid", fgColor="FFFF00")
fill_red = PatternFill("solid", fgColor="C62828")
fill_amber = PatternFill("solid", fgColor="FFF4D6")
thin = Side(style="thin", color="C9D3DE")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")

wb = openpyxl.Workbook()
wb.remove(wb.active)


def sheet(name, title, note=None, widths=None):
    ws = wb.create_sheet(name)
    ws["A1"] = title
    ws["A1"].font = f_title
    if note:
        ws["A2"] = note
        ws["A2"].font = f_note
    for i, w in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False
    return ws


def header(ws, row, labels, col=1):
    for j, lab in enumerate(labels):
        c = ws.cell(row=row, column=col + j, value=lab)
        c.font, c.fill, c.alignment, c.border = f_head, fill_head, Alignment(wrap_text=True, vertical="center"), box


def put(ws, row, col, value, font=f_base, fmt=None, fill=None, comment=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font, c.alignment, c.border = font, wrap, box
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if comment:
        c.comment = Comment(comment, "Paytitude")
    return c


def table(ws, start_row, labels, rows, widths=None):
    header(ws, start_row, labels)
    for i, r in enumerate(rows, start=start_row + 1):
        for j, v in enumerate(r, start=1):
            put(ws, i, j, v)
    return start_row + len(rows)


# ---------------------------------------------------------------- README
ws = sheet("README", "Paytitude · Paytm UPI Growth Challenge, Track B · Round 2 annexure",
           "Settle & Pay: making Paytm the place merchants spend the money they see arrive · 9 October 2026",
           [28, 110])
r = 4
put(ws, r, 1, "Purpose", f_bold); put(ws, r, 2, "Supporting evidence and the live model behind Paytitude_R2.pdf. Every number in the deck traces to a cell here or to a cited source (sheet 'Sources').")
r += 2
put(ws, r, 1, "Evidence-status legend", f_bold); r += 1
for lab, desc in [("Fact (cited)", "Public, sourced fact; see 'Sources' for URL and verification status"),
                  ("Synthetic hypothesis", "Derived from the 50-row synthetic VOC dataset; for hypothesis stress-testing only, not market prevalence"),
                  ("Projection", "Model output from stated assumptions; a projection, not a forecast or commitment"),
                  ("To validate", "Open item to be measured (validation test, Paytm data or pilot)"),
                  ("Validated (n = __)", "Measured with real target users; n stated. None yet: see 'Validation_Log'")]:
    put(ws, r, 1, lab); put(ws, r, 2, desc); r += 1
r += 1
put(ws, r, 1, "Synthetic-data disclosure", f_bold, fill=fill_amber)
put(ws, r, 2, "Verbatim from the VOC workbook's 'Methodology Notes' sheet:", fill=fill_amber); r += 1
for line in C.SYNTHETIC_DISCLOSURE:
    put(ws, r, 2, line, fill=fill_amber); r += 1
r += 1
put(ws, r, 1, "Cell colour convention", f_bold); r += 1
for lab, desc, font, fill in [("Blue text", "Hard-coded input or scenario lever", f_input, None),
                              ("Black text", "Formula", f_base, None),
                              ("Green text", "Link to another sheet", f_link, None),
                              ("Yellow fill", "Key assumption, or a cell to fill in", f_base, fill_key)]:
    put(ws, r, 1, lab, font, fill=fill); put(ws, r, 2, desc); r += 1
r += 1
put(ws, r, 1, "File map", f_bold); r += 1
for lab, desc in [("Paytitude_R2.pdf", "Submission deck: cover, 15 content slides, annexure divider, annexures A1–A8"),
                  ("Paytitude_R2.pptx", "Live version with animations and speaker notes (Morph renders in PowerPoint 2019/365)"),
                  ("Paytitude_R2_Annexure.xlsx", "This workbook")]:
    put(ws, r, 1, lab); put(ws, r, 2, desc); r += 1
r += 1
put(ws, r, 1, "Sheets", f_bold); r += 1
for lab, desc in [("VOC_Synthetic", "All 50 rows of the synthetic VOC dataset, plus derived settlement profile, quadrant, group and contradiction flag"),
                  ("VOC_Real", "Genuine exploratory conversations (notes to be added)"),
                  ("Validation_Log", "One row per tested user (≥20). Summary formulas on top. EMPTY until the test is run"),
                  ("Segment_Rules", "Inclusion rules and group counts, computed live from VOC_Synthetic"),
                  ("Funnel_Model", "Cautious / base / upside funnel (projection)"),
                  ("Share_Metric", "Main metric: Paytm share of segment vendor-UPI value"),
                  ("Unit_Costs", "Cost per incremental MTU and per ₹ of incremental GMV, with and without build"),
                  ("Sensitivity", "Halving the three key assumptions, plus the gap-days-only floor"),
                  ("Pilot_Design", "90-day randomised pilot"),
                  ("Dependencies_Risks", "Dependency register and risk register"),
                  ("Sources", "Every external source with URL, access date and verification status"),
                  ("Change_Log_R1_R2", "What changed since Round 1 and why")]:
    put(ws, r, 1, lab); put(ws, r, 2, desc); r += 1

# ---------------------------------------------------------------- VOC_Synthetic
src = openpyxl.load_workbook(VOC_SRC)["50 VOCs"]
rows = list(src.iter_rows(values_only=True))
hdr, data = list(rows[0]), [list(x) for x in rows[1:]]
ws = wb.create_sheet("VOC_Synthetic")
ws.sheet_view.showGridLines = False
ncol = len(hdr) + 4
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
ws["A1"] = "SYNTHETIC: hypothesis stress-testing, not field research. Counts are not market shares; verbatims are dataset responses, not interviewed merchants."
ws["A1"].font, ws["A1"].fill, ws["A1"].alignment = f_banner, fill_red, Alignment(vertical="center")
ws.row_dimensions[1].height = 24
ws["A2"] = "Source: Paytm_Track_B_VOCs.xlsx, sheet '50 VOCs' (columns A–Z unchanged). Columns AA–AD added by Paytitude."
ws["A2"].font = f_note


def settle_profile(s):
    s = (s or "").lower()
    if s.startswith("no paytm") or "only" in s:
        return "None/partial"
    if "separate" in s:
        return "Next day, separate account"
    if "next" in s:
        return "Next day"
    if "same" in s:
        return "Same day"
    return "Unclear"


quad_of = {v: k for k, ids in C.QUADRANT.items() for v in ids}
group_of = {v: g for g, ids in C.GROUPS.items() for v in ids}
header(ws, 3, hdr + ["Settlement profile (derived)", "Segment quadrant", "Segment group", "Contradiction flag"])
for i, row in enumerate(data, start=4):
    rid = row[0]
    extra = [settle_profile(row[10]), quad_of.get(rid, ""), group_of.get(rid, "Group 3"),
             "Y" if rid in C.CONTRADICTORY else "N"]
    for j, v in enumerate(row + extra, start=1):
        c = put(ws, i, j, v)
        if j > len(hdr):
            c.fill = fill_pale
for j in range(1, ncol + 1):
    ws.column_dimensions[get_column_letter(j)].width = 16 if j not in (6, 18, 19, 24, 25, 26) else 34
ws.freeze_panes = "B4"
VOC_FIRST, VOC_LAST = 4, 3 + len(data)


def rng(col):
    return f"VOC_Synthetic!${col}${VOC_FIRST}:${col}${VOC_LAST}"


SW, OUTG, PROF, QUAD, GRP, CON, FREQ = "T", "L", "AA", "AB", "AC", "AD", "H"

# ---------------------------------------------------------------- VOC_Real
ws = sheet("VOC_Real", "Genuine exploratory merchant conversations",
           "GENUINE, EXPLORATORY, SMALL SAMPLE. Business types as reported by the team; notes and verbatims not yet supplied to the build. Do not treat as validation (that is Validation_Log).",
           [8, 14, 30, 18, 60, 50, 22])
header(ws, 4, ["#", "Date", "Business type", "City", "Key points", "Verbatim (original language)", "Status"])
for i, bt in enumerate(C.REAL_CONVERSATIONS, start=5):
    for j, v in enumerate([i - 4, "", bt, "", "[notes to be added by team]", "[verbatim to be added]",
                           "Reported; notes pending"], start=1):
        put(ws, i, j, v, fill=fill_key if j in (2, 4, 5, 6) else None)

# ---------------------------------------------------------------- Validation_Log
ws = sheet("Validation_Log", "Solution validation log (≥20 target users)",
           "EMPTY: no validation data has been supplied. Fill one row per tested merchant from row 19; summary recalculates; then run fill_validation.py to update slide 10 and re-export the PDF.",
           [10, 12, 12, 18, 10, 18, 10, 14, 20, 10, 30, 10, 10, 10, 10, 18, 40, 12, 10])
cols = ["ID", "Date", "City", "Business type", "Paytm collector (Y/N)", "Settlement profile (Same day / Next day / Separate account)",
        "Vendor UPI payments per week", "Current app", "Method (Concept test / Prototype walkthrough / Real-world trial)",
        "Task completed unaided (Y/N)", "Commitment made (text)", "Named a supplier payment to move (Y/N)",
        "Agreed to follow-up trial (Y/N)", "Enabled instant settlement in session (Y/N)",
        "Used instant settlement in last 30 days (Y/N)", "Top objection (Trust / Habit / Supplier / Staff / Account set-up / Other)",
        "Verbatim", "Evidence type (Observed / Stated)", "Of last 5 payments, would move (0–5)"]
DATA0, DATAN = 19, 78
L = lambda col: f"${col}${DATA0}:${col}${DATAN}"
summary = [
    ("Merchants tested (n)", f"=COUNTA({L('A')})", "0"),
    ("Meets the ≥20 requirement?", f'=IF(B4>=20,"Yes","No: "&(20-B4)&" more needed")', None),
    ("Completed the prototype task unaided", f'=IFERROR(COUNTIF({L("J")},"Y")/B4,"—")', "0%"),
    ("Named a supplier payment they would move", f'=IFERROR(COUNTIF({L("L")},"Y")/B4,"—")', "0%"),
    ("Agreed to a follow-up trial", f'=IFERROR(COUNTIF({L("M")},"Y")/B4,"—")', "0%"),
    ("Enabled instant settlement in session", f'=IFERROR(COUNTIF({L("N")},"Y")/B4,"—")', "0%"),
    ("Made at least one commitment", f'=IFERROR(SUMPRODUCT(--((({L("L")}="Y")+({L("M")}="Y")+({L("N")}="Y"))>0))/B4,"—")', "0%"),
    ("Used instant settlement in last 30 days", f'=IFERROR(COUNTIF({L("O")},"Y")/B4,"—")', "0%"),
    ("Settlement: same day / next day / separate account",
     f'=COUNTIF({L("F")},"Same day")&" / "&COUNTIF({L("F")},"Next day")&" / "&COUNTIF({L("F")},"Separate account")', None),
    ("Average of last 5 payments they would move", f'=IFERROR(AVERAGE({L("S")}),"—")', "0.0"),
    ("Share of answers observed (not stated)", f'=IFERROR(COUNTIF({L("R")},"Observed")/COUNTA({L("R")}),"—")', "0%"),
    ("Top objection counts: Trust / Habit / Supplier / Staff / Account / Other",
     "=" + '&" / "&'.join(f'COUNTIF({L("P")},"{k}")' for k in ["Trust", "Habit", "Supplier", "Staff", "Account set-up", "Other"]), None),
]
for i, (lab, f, fmt) in enumerate(summary, start=4):
    put(ws, i, 1, lab, f_bold)
    ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=4)
    put(ws, i, 5, f, fmt=fmt)
    ws.merge_cells(start_row=i, start_column=5, end_row=i, end_column=7)
# summary formulas reference B4; move n into column E -> fix references
for i in range(4, 4 + len(summary)):
    v = ws.cell(row=i, column=5).value
    ws.cell(row=i, column=5).value = v.replace("B4", "$E$4")
ex = ["EXAMPLE (not data)", "12 Oct 2026", "Pune", "Hardware retailer", "Y", "Next day", 5, "PhonePe",
      "Prototype walkthrough", "Y", "Will pay Monday's cable order via Settle & Pay", "Y", "Y", "N", "N", "Habit",
      "Agar collection se seedha jaata hai toh theek hai", "Observed", 2]
for j, v in enumerate(ex, start=1):
    c = put(ws, 17, j, v, Font(name=FONT, size=9, italic=True, color="8A96A3"))
header(ws, 18, cols)
ws.row_dimensions[18].height = 64
for rr in range(DATA0, DATAN + 1):
    for j in range(1, len(cols) + 1):
        put(ws, rr, j, None, fill=fill_key if rr < DATA0 + 20 else None)
ws.freeze_panes = "B19"

# ---------------------------------------------------------------- Segment_Rules
ws = sheet("Segment_Rules", "Segment rules and group counts (synthetic hypothesis)",
           "All counts are COUNTIFS over VOC_Synthetic. SYNTHETIC: illustrates the hypothesis, not market prevalence.",
           [44, 46, 12, 10, 10, 10, 60])
put(ws, 4, 1, "Group 1 inclusion rules (all must hold)", f_bold)
rules = ["Regular Paytm collector (Paytm QR or Soundbox)", "Next-day settlement, or settlement into a separate current account",
         "Pays vendors by UPI at least weekly", "Payment app chosen by owner or manager",
         "States a timing, transfer or staff-control friction"]
for i, t in enumerate(rules, start=5):
    put(ws, i, 1, f"{i - 4}. {t}")
    ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=2)

header(ws, 11, ["Group", "Rule", "Merchants", "Yes", "Maybe", "No", "Rows"])
grp_rows = [("Group 1", "High potential: all five rules", " ".join(C.GROUPS["Group 1"])),
            ("Group 2", "Paytm collector not in Group 1 or 3: next-day with no friction (7) or same-day with a friction (9, of which 7 contradictory)", " ".join(C.GROUPS["Group 2"])),
            ("Group 3", "No or partial Paytm collection (8); accountant chooses app (2); same-day with no friction (8)", "All other rows")]
for i, (g, rule, ids) in enumerate(grp_rows, start=12):
    put(ws, i, 1, g, f_bold); put(ws, i, 2, rule)
    put(ws, i, 3, f'=COUNTIF({rng(GRP)},A{i})')
    for j, ans in zip((4, 5, 6), ("Yes", "Maybe", "No")):
        put(ws, i, j, f'=COUNTIFS({rng(GRP)},$A{i},{rng(SW)},"{ans}")')
    put(ws, i, 7, ids)
put(ws, 15, 1, "Total", f_bold)
for j in (3, 4, 5, 6):
    col = get_column_letter(j)
    put(ws, 15, j, f"=SUM({col}12:{col}14)", f_bold)

header(ws, 17, ["Settlement profile (Paytm collectors)", "Note", "Merchants", "Yes", "Maybe", "No", ""])
prof = ["Same day", "Next day", "Next day, separate account"]
for i, p in enumerate(prof, start=18):
    put(ws, i, 1, p); put(ws, i, 2, "")
    put(ws, i, 3, f'=COUNTIF({rng(PROF)},A{i})')
    for j, ans in zip((4, 5, 6), ("Yes", "Maybe", "No")):
        put(ws, i, j, f'=COUNTIFS({rng(PROF)},$A{i},{rng(SW)},"{ans}")')
put(ws, 21, 1, "Regular Paytm collectors", f_bold); put(ws, 21, 3, "=SUM(C18:C20)", f_bold)

header(ws, 23, ["Model proxy", "Definition", "Value", "", "", "", "Status"])
proxies = [
    ("Next-day or separate-account merchants", "Next day + next day, separate account", "=C19+C20", "0", "Synthetic hypothesis"),
    ("Settlement share", "Next-day or separate ÷ regular collectors (23 of 42)", "=C24/C21", "0.0%", "Synthetic hypothesis; replace with Paytm settlement data"),
    ("Group 1 merchants", "Next-day or separate with a stated friction", "=C12", "0", "Synthetic hypothesis"),
    ("Friction share", "Group 1 ÷ next-day or separate (16 of 23)", "=C26/C24", "0.0%", "Synthetic hypothesis"),
    ("Group 1 yes", "Stated 'would switch: yes'", "=D12", "0", "Synthetic hypothesis"),
    ("Group 1 maybe", "Stated 'would switch: maybe'", "=E12", "0", "Synthetic hypothesis"),
    ("Group 1 never or rarely pay via Paytm", "Paytm outgoing usage = Never or Rarely",
     f'=COUNTIFS({rng(GRP)},"Group 1",{rng(OUTG)},"Never")+COUNTIFS({rng(GRP)},"Group 1",{rng(OUTG)},"Rarely")', "0", "Synthetic hypothesis"),
    ("Not already a Paytm payer", "Never or rarely ÷ Group 1 (9 of 16)", "=C30/C26", "0.0%", "Synthetic hypothesis"),
    ("Same-day merchants saying 'no'", "Of 19 same-day collectors", "=F18&\" of \"&C18", None, "Synthetic hypothesis"),
    ("Pay vendors by UPI at least 2–3 times a week", "Of 50 rows",
     f'=COUNTA({rng("A")})-COUNTIF({rng(FREQ)},"Weekly")-COUNTIF({rng(FREQ)},"Rarely")', "0", "Synthetic hypothesis"),
    ("Contradictory rows", "Flagged V10, V11, V12, V17, V19, V23, V34, V35, V36, V40, V42", f'=COUNTIF({rng(CON)},"Y")', "0", "Audit"),
    ("Contradictory rows inside Group 1", "Group 1 is unaffected by the clean-sample sensitivity", f'=COUNTIFS({rng(GRP)},"Group 1",{rng(CON)},"Y")', "0", "Audit"),
]
PROXY = {}
for i, (lab, d, f, fmt, st) in enumerate(proxies, start=24):
    put(ws, i, 1, lab, f_bold); put(ws, i, 2, d); put(ws, i, 3, f, fmt=fmt); put(ws, i, 7, st)
    PROXY[lab] = f"Segment_Rules!$C${i}"

header(ws, 37, ["2×2 quadrant (slide 5)", "Rows", "Merchants", "Yes", "Maybe", "No", ""])
for i, (q, ids) in enumerate(C.QUADRANT.items(), start=38):
    put(ws, i, 1, q); put(ws, i, 2, " ".join(ids))
    put(ws, i, 3, f'=COUNTIF({rng(QUAD)},A{i})')
    for j, ans in zip((4, 5, 6), ("Yes", "Maybe", "No")):
        put(ws, i, j, f'=COUNTIFS({rng(QUAD)},$A{i},{rng(SW)},"{ans}")')

# ---------------------------------------------------------------- Funnel_Model
ws = sheet("Funnel_Model", "Incremental growth funnel: cautious / base / upside",
           "PROJECTION, NOT FORECAST. Run-rate at the end of year one. Every rate is an assumption or a synthetic proxy, to be replaced by Paytm data and pilot results.",
           [46, 16, 16, 16, 64])
header(ws, 4, ["Driver", "Cautious", "Base", "Upside", "Source / status"])
FR = {}


def frow(r, label, vals, fmt, status, kind="input", key=None):
    put(ws, r, 1, label, f_bold if kind == "result" else f_base)
    for j, v in enumerate(vals, start=2):
        font = {"input": f_input, "link": f_link, "calc": f_base, "result": f_bold}[kind]
        put(ws, r, j, v, font, fmt=fmt, fill=fill_key if key else None)
    put(ws, r, 5, status)
    FR[label] = r


frow(5, "Paytm device merchants", [15700000] * 3, "#,##0", "Fact: Paytm Q1 FY27 release (S1)")
frow(6, "Share paying vendors by UPI at least weekly", [0.35, 0.5, 0.6], "0.0%", "Assumption (dataset pre-selected for UPI payers)", key=True)
frow(7, "Share settled next day or into a separate account", [f"={PROXY['Settlement share']}"] * 3, "0.0%", "Synthetic proxy: 23 of 42 collectors; replace with Paytm settlement data", "link", key=True)
frow(8, "Share with a stated friction", [f"={PROXY['Friction share']}"] * 3, "0.0%", "Synthetic proxy: 16 of 23", "link")
frow(9, "Eligible merchants", [f"={c}5*{c}6*{c}7*{c}8" for c in "BCD"], "#,##0", "Result", "result")
frow(10, "Share exposed in year one (field force + in-app)", [0.25, 0.4, 0.55], "0.0%", "Assumption")
frow(11, "Exposed merchants", [f"={c}9*{c}10" for c in "BCD"], "#,##0", "Result", "result")
frow(12, "Conversion of stated 'yes'", [0.25, 0.4, 0.6], "0%", "Assumption")
frow(13, "Conversion of stated 'maybe'", [0.05, 0.1, 0.2], "0%", "Assumption")
frow(14, "Group 1 'yes' count", [f"={PROXY['Group 1 yes']}"] * 3, "0", "Synthetic (7 of 16)", "link")
frow(15, "Group 1 'maybe' count", [f"={PROXY['Group 1 maybe']}"] * 3, "0", "Synthetic (9 of 16)", "link")
frow(16, "Group 1 merchants", [f"={PROXY['Group 1 merchants']}"] * 3, "0", "Synthetic", "link")
frow(17, "First-time activation rate", [f"=({c}12*{c}14+{c}13*{c}15)/{c}16" for c in "BCD"], "0.0%", "Assumption applied to dataset propensity", "calc", key=True)
frow(18, "Activated merchants", [f"={c}11*{c}17" for c in "BCD"], "#,##0", "Result", "result")
frow(19, "Repeat rate (2+ payments in each of months 2 and 3)", [0.35, 0.5, 0.65], "0%", "Assumption", key=True)
frow(20, "Repeat merchants", [f"={c}18*{c}19" for c in "BCD"], "#,##0", "Result", "result")
frow(21, "Not already a Paytm payer", [f"={PROXY['Not already a Paytm payer']}"] * 3, "0.0%", "Synthetic proxy: 9 of 16 Group 1 never or rarely pay via Paytm", "link")
frow(22, "Incremental MTU", [f"={c}20*{c}21" for c in "BCD"], "#,##0", "Result", "result")
frow(23, "Weeks per month", [4.3] * 3, "0.0", "Convention")
frow(24, "Vendor UPI payments per merchant per month", [f"=(2.5*{c}23+4*{c}23+5*{c}23+26)/4" for c in "BCD"], "0.00",
     "Synthetic: midpoints of 2–3, 3–5, 4–6 a week and daily (26 a month), equally weighted", "calc")
frow(25, "Share of those payments shifted to Paytm", [0.25, 0.4, 0.6], "0%", "Assumption")
frow(26, "Incremental Paytm business payments a month", [f"={c}20*{c}24*{c}25" for c in "BCD"], "#,##0", "Result", "result")
frow(27, "Average vendor payment (₹)", [2000, 2500, 3500], "#,##0", "Assumption")
frow(28, "Baseline already on Paytm (share of shifted value)", [0.15] * 3, "0%",
     "Assumption. Occasional Paytm payers' shifted payments are assumed mostly new; verify in pilot")
ws.cell(row=28, column=3).comment = Comment("GMV baseline overlap (15%) vs MTU exclusion (44% already sometimes pay via Paytm): occasional Paytm payers' shifted payments are assumed mostly new; verify in pilot.", "Paytitude")
frow(29, "Incremental business UPI GMV a month (₹)", [f"={c}26*{c}27*(1-{c}28)" for c in "BCD"], "#,##0", "Result", "result")
frow(30, "Incremental GMV a month (₹ crore)", [f"={c}29/10^7" for c in "BCD"], "#,##0.0", "Result", "result")
frow(31, "Annual run-rate (₹ crore)", [f"={c}30*12" for c in "BCD"], "#,##0", "Result: deck uses ≈₹240 / ≈₹2,660 / ≈₹19,400 crore (3 significant figures)", "result")
frow(32, "Incremental MTU per 1,000 exposed", [f"={c}22/{c}11*1000" for c in "BCD"], "0.0", "Result", "result")

# ---------------------------------------------------------------- Share_Metric
ws = sheet("Share_Metric", "Main metric: increase in Paytm's share of segment UPI GMV",
           "PROJECTION. The segment pool is the addressable outgoing vendor-UPI value, NOT incremental Paytm GMV.",
           [52, 16, 16, 16, 60])
header(ws, 4, ["Item", "Cautious", "Base", "Upside", "Formula / note"])
srows = [
    ("Eligible merchants", "={F}{c}9", "#,##0", "Funnel_Model row 9", "link"),
    ("Exposed merchants", "={F}{c}11", "#,##0", "Funnel_Model row 11", "link"),
    ("Vendor UPI payments per merchant per month", "={F}{c}24", "0.00", "Funnel_Model row 24", "link"),
    ("Average vendor payment (₹)", "={F}{c}27", "#,##0", "Funnel_Model row 27", "link"),
    ("Segment pool a month (₹ crore)", "={c}5*{c}7*{c}8/10^7", "#,##0", "Eligible × payments × ticket: addressable, not incremental", "calc"),
    ("Segment pool a year (₹ lakh crore)", "={c}9*12/10^5", "0.00", "× 12", "calc"),
    ("Exposed-cohort pool a month (₹ crore)", "={c}6*{c}7*{c}8/10^7", "#,##0", "Exposed × payments × ticket", "calc"),
    ("Incremental Paytm GMV a month (₹ crore)", "={F}{c}30", "#,##0.0", "Funnel_Model row 30", "link"),
    ("Share increase across exposed cohort (pts)", "={c}12/{c}11*100", "0.0", "Incremental ÷ exposed pool", "result"),
    ("Share increase across whole eligible segment (pts)", "={c}12/{c}9*100", "0.0", "Incremental ÷ segment pool, end of year one", "result"),
    ("Expected intention-to-treat lift in pilot (pts, before overlap)", "={F}{c}17*{F}{c}19*{F}{c}25*100", "0.0", "Activation × repeat × share shifted; pilot target +4 to +6 pts", "result"),
    ("Expected intention-to-treat lift in pilot (pts, net of 15% overlap)", "={c}15*(1-{F}{c}28)", "0.0", "As above × (1 − baseline overlap)", "result"),
]
for i, (lab, tpl, fmt, note, kind) in enumerate(srows, start=5):
    put(ws, i, 1, lab, f_bold if kind == "result" else f_base)
    for c in "BCD":
        put(ws, i, "BCD".index(c) + 2, tpl.format(c=c, F="Funnel_Model!"),
            {"link": f_link, "calc": f_base, "result": f_bold}[kind], fmt=fmt)
    put(ws, i, 5, note)
header(ws, 19, ["Early indicator", "Target", "", "", "Status"])
for i, (lab, tgt) in enumerate([("Activation within 30 days of exposure", "15–25%"),
                                 ("Repeat: 2+ payments in each of months 2 and 3", "50%"),
                                 ("Settle & Pay payments per repeat merchant per month", "6–8"),
                                 ("Habit guard: share of payments made when no settlement gap existed", "≥30%"),
                                 ("Incremental MTU per 1,000 exposed", "≈65 (base)"),
                                 ("Wrong-payee payments per 1,000", "Near zero (commitment)")], start=20):
    put(ws, i, 1, lab); put(ws, i, 2, tgt); put(ws, i, 5, "Hypothesis, to validate in pilot")

# ---------------------------------------------------------------- Unit_Costs
ws = sheet("Unit_Costs", "Unit costs: per incremental MTU and per ₹ of incremental GMV (base case)",
           "Every input is a labelled team assumption, not a Paytm figure. No cashback anywhere in the design.",
           [58, 18, 70])
header(ws, 4, ["Item", "Value", "Note"])
urows = [
    ("Field activation visit cost per exposed merchant (₹)", 200, "#,##0", "Assumption: Soundbox field force marginal cost (range ₹150–300)", "input"),
    ("Visit cost, low (₹)", 150, "#,##0", "Assumption", "input"),
    ("Visit cost, high (₹)", 300, "#,##0", "Assumption", "input"),
    ("Float rate (annual)", 0.11, "0%", "Assumption: cost of releasing funds early", "input"),
    ("Days settled early", 1, "0", "Assumption", "input"),
    ("Average vendor payment (₹)", "=Funnel_Model!C27", "#,##0", "Link", "link"),
    ("Free early settlement cost per payment (₹)", "=B10*B8/365*B9", "0.00", "Ceiling: settling already-received funds mainly forgoes float income", "calc"),
    ("In-app prompts, notifications, WhatsApp (₹ per merchant)", 0, "0", "Assumption: ≈₹0 marginal", "input"),
    ("Pilot build: prompt, hand-off, fee waiver (₹)", 5000000, "#,##0", "Assumption: small build ≈₹0.5 crore, one-time", "input"),
    ("Full in-app payer flow (₹)", 50000000, "#,##0", "Assumption: medium-large build ≈₹5 crore, only after the scale gate", "input"),
    ("Exposed merchants", "=Funnel_Model!C11", "#,##0", "Link", "link"),
    ("Incremental MTU", "=Funnel_Model!C22", "#,##0", "Link", "link"),
    ("Incremental payments a year", "=Funnel_Model!C26*12", "#,##0", "Run-rate × 12", "link"),
    ("Incremental GMV a year (₹)", "=Funnel_Model!C29*12", "#,##0", "Annual run-rate", "link"),
    ("Field visit cost (₹)", "=B5*B15", "#,##0", "Visit cost × exposed", "calc"),
    ("Messaging cost (₹)", "=B12*B15", "#,##0", "", "calc"),
    ("Free settlement cost a year (₹)", "=B11*B17", "#,##0", "Cost per payment × payments", "calc"),
    ("Total year-one cost, without build (₹)", "=B19+B20+B21", "#,##0", "", "result"),
    ("Total year-one cost, with build (₹)", "=B22+B13+B14", "#,##0", "", "result"),
    ("Cost per incremental MTU, without build (₹)", "=B22/B16", "#,##0", "", "result"),
    ("Cost per incremental MTU, with build (₹)", "=B23/B16", "#,##0", "", "result"),
    ("Cost per ₹ of incremental GMV, without build (₹)", "=B22/B18", "0.0000", "Year-one cost ÷ annual run-rate GMV", "result"),
    ("Cost per ₹ of incremental GMV, with build (₹)", "=B23/B18", "0.0000", "", "result"),
    ("Cost per ₹100 of incremental GMV, without build (₹)", "=B26*100", "0.00", "", "result"),
    ("Cost per ₹100 of incremental GMV, with build (₹)", "=B27*100", "0.00", "", "result"),
    ("Ramp-adjusted cost per ₹, with build (year-one GMV = 50% of run-rate)", "=B23/(B18*0.5)", "0.0000", "Sensitivity on the denominator", "result"),
    ("Pilot metric: visit cost per incremental MTU (₹200 × 1,000 ÷ MTU per 1,000)", "=B5*1000/Funnel_Model!C32", "#,##0", "", "result"),
    ("Cost per incremental MTU at visit cost ₹150, without build (₹)", "=(B6*B15+B20+B21)/B16", "#,##0", "", "result"),
    ("Cost per incremental MTU at visit cost ₹300, without build (₹)", "=(B7*B15+B20+B21)/B16", "#,##0", "", "result"),
]
for i, (lab, v, fmt, note, kind) in enumerate(urows, start=5):
    put(ws, i, 1, lab, f_bold if kind == "result" else f_base)
    put(ws, i, 2, v, {"input": f_input, "link": f_link, "calc": f_base, "result": f_bold}[kind], fmt=fmt,
        fill=fill_key if lab.startswith(("Field activation", "Pilot build", "Full in-app")) else None)
    put(ws, i, 3, note)
r = 5 + len(urows) + 1
put(ws, r, 1, "Durability once incentives stop", f_bold); r += 1
for t in ["No cashback in the design. The only economic element is free settlement when it funds a supplier payment: permanent product pricing, not a promotion.",
          "Durability test (a): repeat use after field-activation contact ends (days 31–90).",
          "Durability test (b): habit-guard share, i.e. usage when no settlement benefit applies.",
          "Durability test (c): optional sub-arm where free settlement reverts to standard pricing for days 61–90, to measure price sensitivity."]:
    put(ws, r, 1, t); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3); r += 1

# ---------------------------------------------------------------- Sensitivity
ws = sheet("Sensitivity", "Sensitivity: the three assumptions that move the answer most (base case)",
           "PROJECTION. Each row recomputes the full chain with one driver changed. The model is multiplicative, so halving any of the three halves every output.",
           [36, 14, 14, 14, 16, 16, 16, 18, 12])
header(ws, 4, ["Scenario", "Settlement share", "Activation", "Repeat", "Shifted-payment factor", "Eligible", "Incremental MTU", "Annual GMV (₹ crore)", "vs base"])
put(ws, 3, 1, "Habit-guard target (bound for the gap-days floor)", f_bold); put(ws, 3, 5, 0.3, f_input, fmt="0%", fill=fill_key)
scen = [("Base case", "=Funnel_Model!$C$7", "=Funnel_Model!$C$17", "=Funnel_Model!$C$19", 1),
        ("Settlement share halved", "=Funnel_Model!$C$7/2", "=Funnel_Model!$C$17", "=Funnel_Model!$C$19", 1),
        ("Activation halved", "=Funnel_Model!$C$7", "=Funnel_Model!$C$17/2", "=Funnel_Model!$C$19", 1),
        ("Repeat rate halved", "=Funnel_Model!$C$7", "=Funnel_Model!$C$17", "=Funnel_Model!$C$19/2", 1),
        ("Gap-days-only floor (base × (1 − habit guard))", "=Funnel_Model!$C$7", "=Funnel_Model!$C$17", "=Funnel_Model!$C$19", "=1-$E$3"),
        ("Share of payments shifted halved (GMV only)", "=Funnel_Model!$C$7", "=Funnel_Model!$C$17", "=Funnel_Model!$C$19", 0.5)]
for i, (lab, s, a, rp, g) in enumerate(scen, start=5):
    put(ws, i, 1, lab, f_bold)
    put(ws, i, 2, s, f_link, fmt="0.0%"); put(ws, i, 3, a, f_link, fmt="0.0%"); put(ws, i, 4, rp, f_link, fmt="0%")
    put(ws, i, 5, g, f_input if isinstance(g, (int, float)) else f_base, fmt="0%")
    put(ws, i, 6, f"=Funnel_Model!$C$5*Funnel_Model!$C$6*B{i}*Funnel_Model!$C$8", fmt="#,##0")
    put(ws, i, 7, f"=F{i}*Funnel_Model!$C$10*C{i}*D{i}*Funnel_Model!$C$21", fmt="#,##0")
    put(ws, i, 8, f"=F{i}*Funnel_Model!$C$10*C{i}*D{i}*Funnel_Model!$C$24*Funnel_Model!$C$25*Funnel_Model!$C$27*(1-Funnel_Model!$C$28)*E{i}*12/10^7", fmt="#,##0")
    put(ws, i, 9, f"=H{i}/$H$5", fmt="0%")
put(ws, 12, 1, "Note: the gap-days-only floor assumes only payments made when a settlement gap exists move to Paytm; the pilot's habit-guard target (≥30% of payments with no gap) bounds the complement. MTU is unchanged in the floor and shift-share rows because the same merchants stay active.", f_note)
ws.merge_cells("A12:I12")
ws.row_dimensions[12].height = 40

# ---------------------------------------------------------------- Pilot_Design
ws = sheet("Pilot_Design", "90-day randomised pilot", "Design, metrics and decision gates", [26, 120])
table(ws, 4, ["Element", "Content"], C.PILOT)

# ---------------------------------------------------------------- Dependencies_Risks
ws = sheet("Dependencies_Risks", "Dependency register and risk register", "Owners are proposed functions, not confirmed people.",
           [36, 60, 50, 26, 18])
end = table(ws, 4, ["Mechanic / dependency", "What public sources indicate", "Still to confirm", "Owner", "Build effort"], C.DEPENDENCIES)
table(ws, end + 3, ["Risk", "Signal", "Mitigation"], C.RISKS)

# ---------------------------------------------------------------- Sources
ws = sheet("Sources", "Sources", f"Access date for web sources: {C.ACCESS_DATE}. Verification status as of build.", [6, 50, 60, 60, 40])
table(ws, 4, ["ID", "Source", "URL", "Used for", "Verification status"], C.SOURCES)

# ---------------------------------------------------------------- Change_Log
ws = sheet("Change_Log_R1_R2", "What changed since Round 1, and why",
           "Round 1 column verified against Paytitude_ExecSummary.pdf (25 Sep 2026). Track B kept.", [20, 55, 55, 50])
table(ws, 4, ["Area", "Round 1", "Round 2", "Why"], C.CHANGE_LOG)

for w in wb.worksheets:
    w.sheet_properties.tabColor = {"VOC_Synthetic": "C62828", "Validation_Log": "FFB300"}.get(w.title, NAVY)
wb.save(OUT)
print("saved", OUT)
